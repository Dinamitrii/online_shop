"""Printable order snapshots and portable Word/Excel/CSV downloads."""
import csv
import io
import math
import re
from decimal import Decimal, ROUND_HALF_UP



_XML_INVALID = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


def clean(value):
    return _XML_INVALID.sub('', str(value if value is not None else ''))


def money(value):
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def snapshot(order, shipment=None):
    meta = [('Номер на поръчката', str(order.id)),
            ('Дата', order.created_at.strftime('%d.%m.%Y %H:%M')),
            ('Статус', order.status), ('Клиент', order.customer_name),
            ('Телефон', order.phone), ('Имейл', order.email or ''),
            ('Адрес от поръчката', order.address)]
    if shipment:
        meta.append(('Товарителница Еконт', shipment.shipment_number))
        meta.append(('Среда на Еконт', 'Тестова' if shipment.environment == 'test' else 'Реална'))
    return {'id': order.id, 'title': f'Документ за поръчка {order.id}',
            'metadata': [(k, clean(v)) for k, v in meta],
            'items': [{'name': clean(i.product_name), 'qty': i.qty,
                       'price': money(i.price), 'subtotal': money(i.subtotal)}
                      for i in sorted(order.items, key=lambda i: i.id)],
            'total': money(order.total)}


def csv_bytes(data):
    def cell(value):
        value = clean(value)
        if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
            return "'" + value
        return value
    output = io.StringIO(newline='')
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_ALL)
    meta = data['metadata']
    writer.writerow([k for k, _ in meta] + ['Артикул', 'Количество', 'Ед. цена EUR', 'Сума EUR', 'Общо поръчка EUR'])
    for item in data['items'] or [None]:
        values = [v for _, v in meta] + ([item['name'], item['qty'], f"{item['price']:.2f}",
                    f"{item['subtotal']:.2f}"] if item else ['', '', '', '']) + [f"{data['total']:.2f}"]
        writer.writerow([cell(v) for v in values])
    return output.getvalue().encode('utf-8-sig')


def docx_bytes(data):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

    doc = Document()
    # Remove decorative title rules inherited from the bundled Word template.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    for name in ('Normal', 'Title', 'Heading 1'):
        style = doc.styles[name]
        style.font.name = 'DejaVu Sans'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    doc.add_paragraph('Железария Дианабад')
    doc.add_paragraph(data['title'], 'Title')
    for key, value in data['metadata']:
        if value:
            p = doc.add_paragraph()
            p.add_run(key + ': ').bold = True
            p.add_run(value)
    doc.add_paragraph('Поръчани артикули', 'Heading 1')
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.9, 8.1, 2.0, 3.2, 3.2]
    for column, width in zip(table.columns, widths):column.width = Cm(width)
    headers = ['№', 'Артикул', 'Брой', 'Ед. цена EUR', 'Сума EUR']
    for c, text in zip(table.rows[0].cells, headers):c.text = text
    repeat = OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, item in enumerate(data['items'], 1):
        cells = table.add_row().cells
        for c, text in zip(cells, [str(i), item['name'], str(item['qty']),
                                  f"{item['price']:.2f}", f"{item['subtotal']:.2f}"]):c.text = text
    if not data['items']:table.add_row().cells[1].text = 'Няма артикули'
    for index, row in enumerate(table.rows):
        no_split = OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for col, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement('w:' + side)
                for k, v in [('val','single'),('sz','4'),('color','D9D9D9')]:edge.set(qn('w:'+k),v)
                borders.append(edge)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'90');edge.set(qn('w:type'),'dxa');margins.append(edge)
            props.append(margins)
            if index == 0 or index % 2 == 0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2' if index==0 else 'F7F8FA');props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col == 1 else WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs:run.font.size=Pt(9);run.bold=(index==0)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.add_run(f"Общо за артикулите: {data['total']:.2f} EUR").bold = True
    doc.add_paragraph('Доставката се начислява отделно от куриера.')
    doc.core_properties.title = data['title']
    doc.core_properties.author = 'Железария Дианабад'
    output=io.BytesIO();doc.save(output);return output.getvalue()


def xlsx_bytes(data):
    # Portable server-side exporter: production Flask does not depend on Codex tools.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.utils import get_column_letter

    wb=Workbook();ws=wb.active;ws.title='Поръчка'
    def text(row, col, value):
        cell=ws.cell(row,col,clean(value));cell.data_type='s';cell.number_format='@';return cell
    ws.sheet_view.showGridLines=False
    ws.merge_cells('A1:E1');text(1,1,'Железария Дианабад').font=Font(size=12,bold=True)
    ws.merge_cells('A2:E2');text(2,1,data['title']).font=Font(size=18,bold=True)
    ws.row_dimensions[2].height=30
    row=4
    for key,value in data['metadata']:
        ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=5)
        text(row,1,key).font=Font(bold=True,size=10)
        text(row,2,value).alignment=Alignment(wrap_text=True,vertical='center',horizontal='left')
        ws.row_dimensions[row].height=max(30,15*math.ceil(len(value)/70))
        row+=1
    row+=1;header=row
    for col,name in enumerate(['№','Артикул','Количество','Ед. цена EUR','Сума EUR'],1):
        text(row,col,name)
    for number,item in enumerate(data['items'],1):
        row+=1
        ws.cell(row,1,number);text(row,2,item['name']);ws.cell(row,3,item['qty'])
        ws.cell(row,4,float(item['price']));ws.cell(row,5,float(item['subtotal']))
        ws.row_dimensions[row].height=max(30,16*math.ceil(len(item['name'])/42))
    last=row
    row+=2;text(row,2,'Общо за артикулите EUR').font=Font(bold=True)
    ws.cell(row,5,float(data['total'])).font=Font(bold=True)
    ws.cell(row,5).number_format='#,##0.00'
    row+=2;ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=5)
    text(row,1,'Доставката се начислява отделно от куриера.')
    border=Border(*(Side(style='thin',color='D9D9D9') for _ in range(4)))
    for cells in ws.iter_rows(min_row=header,max_row=last,max_col=5):
        for cell in cells:
            cell.border=border
            cell.alignment=Alignment(wrap_text=True,vertical='center',horizontal='left' if cell.column==2 else 'right')
            if cell.row==header:
                cell.fill=PatternFill('solid',fgColor='17324D');cell.font=Font(color='FFFFFF',bold=True)
            elif (cell.row-header)%2==0:cell.fill=PatternFill('solid',fgColor='F2F5F8')
            if cell.row>header and cell.column in (4,5):cell.number_format='#,##0.00'
    for i,width in enumerate([23,46,14,19,19],1):ws.column_dimensions[get_column_letter(i)].width=width
    ws.row_dimensions[header].height=30
    for cells in ws.iter_rows(min_row=4,max_row=header-2,max_col=1):
        cells[0].alignment=Alignment(wrap_text=True,vertical='center')
    ws.freeze_panes=f'C{header+1}'
    if last>header:ws.auto_filter.ref=f'A{header}:E{last}'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True
    ws.page_margins=PageMargins(left=.3,right=.3,top=.4,bottom=.4,header=.2,footer=.2)
    ws.print_title_rows=f'{header}:{header}';ws.print_area=f'A1:E{row}'
    ws.oddFooter.center.text='Страница &P от &N'
    output=io.BytesIO();wb.save(output);return output.getvalue()


def register_order_documents(app, admin_required):
    from flask import Response, abort, render_template
    from models import Order, CourierShipment

    def data_for(order_id):
        order=Order.query.get_or_404(order_id)
        shipment=CourierShipment.query.filter_by(order_id=order.id,
            environment=app.config['COURIER_ENVIRONMENT'],state='created').first()
        return snapshot(order,shipment)

    @app.route('/admin/orders/<int:order_id>/print')
    @admin_required
    def order_print(order_id):
        response=Response(render_template('admin/order_print.html', document=data_for(order_id)))
        response.headers['Cache-Control']='no-store'
        return response

    @app.route('/admin/orders/<int:order_id>/document.<extension>')
    @admin_required
    def order_document(order_id,extension):
        exporters={'docx':(docx_bytes,'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
                   'xlsx':(xlsx_bytes,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                   'csv':(csv_bytes,'text/csv; charset=utf-8')}
        if extension not in exporters:abort(404)
        builder,mime=exporters[extension]
        try:
            output=builder(data_for(order_id))
        except ImportError:
            return 'Библиотеките за Word и Excel не са инсталирани. Обновете зависимостите от requirements.txt.',503
        response=Response(output,content_type=mime)
        response.headers['Content-Disposition']=f'attachment; filename="order-{order_id}.{extension}"'
        response.headers['Cache-Control']='no-store'
        response.headers['X-Content-Type-Options']='nosniff'
        return response
