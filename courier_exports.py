"""Local exports of the saved shipment; the official shipping label remains Econt's PDF."""
import csv
import io
import json
import re

NOTE = 'Справка по запазените данни за пратката. За предаване на куриера използвайте официалния PDF от Еконт.'


def clean(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value if value is not None else ''))


def shipment_rows(shipment):
    label = json.loads(shipment.request_json)
    sender = label.get('senderClient') or {}
    receiver = label.get('receiverClient') or {}
    address = label.get('receiverAddress') or {}
    city = address.get('city') or {}
    services = label.get('services') or {}
    return [
        ('Товарителница', shipment.shipment_number),
        ('Поръчка', shipment.order_id),
        ('Среда', 'ТЕСТОВА' if shipment.environment == 'test' else 'Реална'),
        ('Статус', shipment.delivery_status),
        ('Подател', sender.get('name', '')),
        ('Телефон на подателя', ', '.join(sender.get('phones') or [])),
        ('Офис на изпращане', label.get('senderOfficeCode', '')),
        ('Получател', receiver.get('name', '')),
        ('Телефон на получателя', ', '.join(receiver.get('phones') or [])),
        ('Офис на получаване', label.get('receiverOfficeCode', '')),
        ('Населено място', city.get('name', '')),
        ('Пощенски код', city.get('postCode', '')),
        ('Адрес', address.get('fullAddress', '')),
        ('Съдържание', label.get('shipmentDescription', '')),
        ('Тегло kg', label.get('weight', '')),
        ('Брой пакети', label.get('packCount', '')),
        ('Наложен платеж', services.get('cdAmount', '')),
        ('Валута на наложения платеж', services.get('cdCurrency', '')),
        ('Доставка за сметка на', 'Получател' if label.get('paymentReceiverAmount') == 100 and label.get('paymentReceiverAmountIsPercent') else 'Подател'),
    ]


def export_shipment(shipment, fmt):
    rows = shipment_rows(shipment)
    output = io.BytesIO()
    if fmt == 'csv':
        text = io.StringIO(newline='')
        writer = csv.writer(text, delimiter=';')
        def safe(value):
            value = clean(value)
            return "'" + value if value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')) else value
        writer.writerow([key for key, _ in rows] + ['Бележка'])
        writer.writerow([safe(value) for _, value in rows] + [NOTE])
        output.write(text.getvalue().encode('utf-8-sig'))
        mime = 'text/csv; charset=utf-8'
    elif fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = 'Пратка'
        ws.append(['Данни за пратката', 'Стойност'])
        for key, value in rows + [('Бележка', NOTE)]:
            ws.append([key, value if isinstance(value, (int, float)) else clean(value)])
            cell = ws.cell(ws.max_row, 2)
            if isinstance(value, str):
                cell.data_type = 's'
                cell.number_format = '@'
            if key in ('Тегло kg', 'Наложен платеж'):
                cell.number_format = '0.00#'
        ws.column_dimensions['A'].width = 34
        ws.column_dimensions['B'].width = 75
        ws.freeze_panes = 'B2'
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name='Calibri', size=11)
            ws.row_dimensions[row[0].row].height = max(30, 16 * (1 + len(clean(row[1].value)) // 65))
        for cell in ws[1]:
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF', size=12)
            cell.fill = PatternFill('solid', fgColor='253D4A')
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:1'
        wb.save(output)
        mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif fmt == 'docx':
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        doc = Document()
        from docx.oxml.ns import qn
        for document_style in doc.styles:
            for border in list(document_style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
        section = doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.6)
        section.left_margin = section.right_margin = Cm(2)
        style = doc.styles['Normal']
        style.font.name, style.font.size = 'Calibri', Pt(10)
        style.paragraph_format.space_after = Pt(3)
        doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
        doc.add_paragraph('Данни за пратка Еконт', 'Title')
        doc.add_paragraph(NOTE)
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Light Shading Accent 1'
        table.autofit = False
        table.columns[0].width, table.columns[1].width = Cm(6), Cm(11)
        for key, value in rows:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = key, clean(value)
            cells[0].paragraphs[0].runs[0].bold = True
        doc.save(output)
        mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    else:
        raise ValueError('Unsupported export format')
    output.seek(0)
    return output, mime
