import csv
import io
import json
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch
from docx import Document
from openpyxl import load_workbook
from test_courier import CourierTests
from courier_exports import export_shipment


class ExportRoutes(CourierTests):
    def test_exports_and_access(self):
        from app import app
        from models import CourierShipment, db
        with patch('courier.EcontClient.label', return_value=self.status):
            self.c.post(self.url, data=self.form)
        for fmt in ('docx', 'xlsx', 'csv'):
            url = self.url + '/export/' + fmt
            response = self.c.get(url)
            self.assertEqual(response.status_code, 200)
            self.assertIn('attachment', response.headers['Content-Disposition'])
            self.assertEqual(response.headers['Cache-Control'], 'no-store')
            self.assertEqual(app.test_client().get(url).status_code, 302)
        self.assertEqual(self.c.get(self.url+'/export/exe').status_code, 404)
        shipment = CourierShipment.query.one()
        shipment.environment = 'live'; db.session.commit()
        self.assertEqual(self.c.get(self.url+'/export/csv').status_code, 404)
        shipment.environment = 'test'; shipment.state = 'failed'; db.session.commit()
        self.assertEqual(self.c.get(self.url+'/export/csv').status_code, 404)


class ExportContent(TestCase):
    def test_content_types_and_spreadsheet_safety(self):
        shipment=SimpleNamespace(shipment_number='00123', order_id=42,environment='test',
            delivery_status='Подготвена',request_json=json.dumps({
                'receiverClient':{'name':'=1+1','phones':['0888888888']},
                'receiverOfficeCode':'0010','weight':1.5,'packCount':2,
                'services':{'cdAmount':12.35,'cdCurrency':'EUR'}}))
        data,_=export_shipment(shipment,'xlsx')
        ws=load_workbook(data).active
        values={r[0].value:r[1] for r in ws.iter_rows(min_row=2)}
        self.assertEqual(values['Получател'].data_type,'s')
        self.assertEqual(values['Получател'].value,'=1+1')
        self.assertEqual(values['Телефон на получателя'].value,'0888888888')
        self.assertEqual(values['Наложен платеж'].value,12.35)
        data,_=export_shipment(shipment,'csv')
        row=next(csv.DictReader(io.StringIO(data.getvalue().decode('utf-8-sig')),delimiter=';'))
        self.assertEqual(row['Получател'],"'=1+1")
        self.assertEqual(row['Телефон на получателя'],'0888888888')
        data,_=export_shipment(shipment,'docx')
        doc=Document(data)
        self.assertIn('00123',' '.join(c.text for r in doc.tables[0].rows for c in r.cells))
