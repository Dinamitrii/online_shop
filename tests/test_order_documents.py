import csv
import io
import unittest
from unittest.mock import patch, MagicMock

import test_courier
from test_courier import app, db, Order, OrderItem, CourierShipment
from docx import Document
from openpyxl import load_workbook
from econt import EcontError, fetch_pdf


class OrderDocumentTests(unittest.TestCase):
    setUp = test_courier.CourierTests.setUp
    tearDown = test_courier.CourierTests.tearDown

    def test_downloads_are_private(self):
        for suffix in ['print','document.docx','document.xlsx','document.csv','econt/print','econt/print.pdf']:
            self.assertEqual(app.test_client().get(f'/admin/orders/{self.order_id}/{suffix}').status_code,302)
        self.assertEqual(self.c.get(f'/admin/orders/{self.order_id}/document.exe').status_code,404)

    def test_docx_contains_order_and_items(self):
        r=self.c.get(f'/admin/orders/{self.order_id}/document.docx')
        self.assertEqual(r.status_code,200)
        self.assertEqual(r.headers['Cache-Control'],'no-store')
        doc=Document(io.BytesIO(r.data))
        paragraphs='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Тест Получател',paragraphs)
        self.assertIn('12.35 EUR',paragraphs)
        self.assertEqual(doc.tables[0].rows[1].cells[1].text,'Болт')
        self.assertEqual(doc.tables[0].rows[1].cells[2].text,'2')

    def test_xlsx_types_and_no_formula_injection(self):
        OrderItem.query.one().product_name='=HYPERLINK("https://evil.test")'
        Order.query.one().customer_name='=1+1';db.session.commit()
        r=self.c.get(f'/admin/orders/{self.order_id}/document.xlsx')
        self.assertEqual(r.status_code,200)
        ws=load_workbook(io.BytesIO(r.data)).active
        values=[cell for row in ws for cell in row]
        self.assertFalse(any(c.data_type=='f' for c in values))
        phone=next(c for c in values if c.value=='0888888888')
        self.assertEqual(phone.data_type,'s')
        self.assertTrue(any(c.data_type=='n' and c.value==12.35 for c in values))
        self.assertTrue(ws.freeze_panes);self.assertTrue(ws.print_area)

    def test_csv_bom_cyrillic_injection_and_quoted_delimiters(self):
        OrderItem.query.one().product_name='=1+1; "болт"\nнов ред';db.session.commit()
        r=self.c.get(f'/admin/orders/{self.order_id}/document.csv')
        self.assertTrue(r.data.startswith(b'\xef\xbb\xbf'))
        rows=list(csv.reader(io.StringIO(r.data.decode('utf-8-sig')),delimiter=';'))
        self.assertEqual(len(rows),2)
        self.assertEqual(rows[1][7], '\'=1+1; "болт"\nнов ред')
        self.assertEqual(rows[1][-1],'12.35')

    def test_print_document_escapes_html_and_has_items(self):
        Order.query.one().customer_name='<script>alert(1)</script>';db.session.commit()
        r=self.c.get(f'/admin/orders/{self.order_id}/print')
        self.assertEqual(r.status_code,200)
        self.assertIn('&lt;script&gt;',r.text)
        self.assertIn('Болт',r.text)
        self.assertIn('window.print()',r.text)
        self.assertIn('12.35',r.text)

    def test_empty_order_exports(self):
        OrderItem.query.delete();db.session.commit()
        for ext in ['docx','xlsx','csv']:
            self.assertEqual(self.c.get(f'/admin/orders/{self.order_id}/document.{ext}').status_code,200)

    def test_waybill_print_original_pdf_without_new_label(self):
        with patch('courier.EcontClient.label',return_value=self.status):self.c.post(self.url,data=self.form)
        with patch('courier.fetch_pdf',return_value=b'%PDF-1.4\nfixture') as pdf, patch('courier.EcontClient.label') as create:
            page=self.c.get(self.url+'/print')
            self.assertEqual(page.status_code,200)
            r=self.c.get(self.url+'/print.pdf')
            self.assertEqual(r.mimetype,'application/pdf')
            self.assertTrue(r.headers['Content-Disposition'].startswith('inline;'))
            self.assertEqual(r.data,b'%PDF-1.4\nfixture')
            pdf.assert_called_once();create.assert_not_called()
        with patch('courier.fetch_pdf',side_effect=EcontError('PDF unavailable')):
            self.assertEqual(self.c.get(self.url+'/print.pdf').status_code,502)
        app.config['COURIER_ENVIRONMENT']='live'
        self.assertEqual(self.c.get(self.url+'/print').status_code,404)

    def test_pdf_fetch_validation_and_no_credentials(self):
        response=MagicMock();response.__enter__.return_value.read.return_value=b'%PDF-1.4\nfixture'
        with patch('econt.build_opener') as builder:
            builder.return_value.open.return_value=response
            self.assertTrue(fetch_pdf('https://demo.econt.com/doc.pdf','test').startswith(b'%PDF-'))
            req=builder.return_value.open.call_args.args[0]
            self.assertNotIn('Authorization',req.headers)
            with self.assertRaises(EcontError):fetch_pdf('https://evil.test/file','test')
            response.__enter__.return_value.read.return_value=b'<html>Error</html>'
            with self.assertRaises(EcontError):fetch_pdf('https://demo.econt.com/doc.pdf','test')

if __name__=='__main__':unittest.main()
